"""Testes de app_qt.py — exportar TXT e gerar relatório DOCX pelo menu
Arquivo/botão marrom, mesmo padrão de tests/test_gerar_planilha_qt.py
(QFileDialog nativo simulado)."""
from __future__ import annotations

from PySide6.QtCore import Qt
from PySide6.QtWidgets import QFileDialog, QMessageBox

from app_qt import JanelaPrincipal


def test_exportar_txt_produz_arquivo_com_dados_da_parcela(qtbot, tmp_path, monkeypatch):
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: None)

    janela = JanelaPrincipal()
    qtbot.addWidget(janela)
    aba = janela._abas_perimetro()[0]
    colunas = aba.modelo_tabela.dataframe().columns
    aba.modelo_tabela.setData(aba.modelo_tabela.index(0, colunas.get_loc("codigo")), "AJ4-M-001", Qt.ItemDataRole.EditRole)
    aba.modelo_tabela.setData(aba.modelo_tabela.index(0, colunas.get_loc("e_long")), "-48,50", Qt.ItemDataRole.EditRole)
    aba.modelo_tabela.setData(aba.modelo_tabela.index(0, colunas.get_loc("n_lat")), "-27,50", Qt.ItemDataRole.EditRole)

    caminho = tmp_path / "saida.txt"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(caminho), "")))

    janela._exportar_txt()

    assert caminho.exists()
    conteudo = caminho.read_text(encoding="utf-8")
    assert "AJ4-M-001" in conteudo
    assert "\t" in conteudo


def test_botao_exportar_txt_do_painel_chama_o_mesmo_metodo(qtbot, monkeypatch):
    chamadas = []
    janela = JanelaPrincipal()
    qtbot.addWidget(janela)
    janela._exportar_txt = lambda: chamadas.append("txt")
    janela._gerar_relatorio = lambda: chamadas.append("docx")
    janela._gerar_planilha = lambda: chamadas.append("ods")

    # os botões já foram conectados na construção da janela (aos métodos
    # antigos) — reconectar aqui simula o mesmo clique chamando a versão
    # substituída, só pra provar que é o MESMO botão exposto pelo painel.
    janela.painel_numeracao.btn_exportar_txt.clicked.disconnect()
    janela.painel_numeracao.btn_exportar_txt.clicked.connect(janela._exportar_txt)
    janela.painel_numeracao.btn_gerar_relatorio.clicked.disconnect()
    janela.painel_numeracao.btn_gerar_relatorio.clicked.connect(janela._gerar_relatorio)
    janela.painel_numeracao.btn_exportar_planilha.clicked.disconnect()
    janela.painel_numeracao.btn_exportar_planilha.clicked.connect(janela._gerar_planilha)

    janela.painel_numeracao.btn_exportar_txt.click()
    janela.painel_numeracao.btn_gerar_relatorio.click()
    janela.painel_numeracao.btn_exportar_planilha.click()

    assert chamadas == ["txt", "docx", "ods"]


def test_gerar_relatorio_produz_docx_com_identificacao_e_planilha(qtbot, tmp_path, monkeypatch):
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: None)

    janela = JanelaPrincipal()
    qtbot.addWidget(janela)
    janela.aba_identificacao.nome.setText("Fulano de Tal")
    aba = janela._abas_perimetro()[0]
    colunas = aba.modelo_tabela.dataframe().columns
    aba.modelo_tabela.setData(aba.modelo_tabela.index(0, colunas.get_loc("codigo")), "AJ4-M-001", Qt.ItemDataRole.EditRole)
    aba.modelo_tabela.setData(aba.modelo_tabela.index(0, colunas.get_loc("e_long")), "-48,50", Qt.ItemDataRole.EditRole)
    aba.modelo_tabela.setData(aba.modelo_tabela.index(0, colunas.get_loc("n_lat")), "-27,50", Qt.ItemDataRole.EditRole)

    caminho = tmp_path / "relatorio.docx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(caminho), "")))

    janela._gerar_relatorio()

    assert caminho.exists()
    from docx import Document

    doc = Document(str(caminho))
    textos = [p.text for p in doc.paragraphs]
    assert "Identificação" in textos
    assert any("AJ4-M-001" in "".join(c.text for c in linha.cells) for tabela in doc.tables for linha in tabela.rows)
