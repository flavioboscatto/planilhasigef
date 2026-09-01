"""Testes de core/relatorio_docx.py — ordem pedida pelo usuário:
Identificação uma vez no topo; por parcela, resumo do cabeçalho,
Importações e por fim a Planilha (só colunas oficiais)."""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from core.models import Identificacao, Perimetro, Vertice
from core.relatorio_docx import gerar_relatorio

_CASAS = {"segundos_gms_casas_decimais": 3, "sigma_casas_decimais": 2, "h_casas_decimais": 3}


def _paragrafos(saida: bytes) -> list[str]:
    doc = Document(io.BytesIO(saida))
    return [p.text for p in doc.paragraphs]


def test_identificacao_aparece_uma_vez_no_topo_antes_das_parcelas():
    identificacao = Identificacao(nome="Fulano de Tal")
    parcelas = [
        {"perimetro": Perimetro(denominacao="Parte 1", vertices=[]), "importacoes": []},
        {"perimetro": Perimetro(denominacao="Parte 2", vertices=[]), "importacoes": []},
    ]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)
    titulos = _paragrafos(saida)

    assert titulos.count("Identificação") == 1
    indice_ident = titulos.index("Identificação")
    # só "Parcela N" — sem "— Denominação" (pedido explícito do usuário)
    indice_p1 = titulos.index("Parcela 1")
    indice_p2 = titulos.index("Parcela 2")
    assert indice_ident < indice_p1 < indice_p2


def test_ordem_por_parcela_e_resumo_importacoes_planilha():
    identificacao = Identificacao()
    parcelas = [
        {
            "perimetro": Perimetro(denominacao="Parte 1", vertices=[Vertice(codigo="A", e_long=-48.0, n_lat=-9.0)]),
            "importacoes": [{"titulo": "Importação 1", "colunas": ["codigo"], "linhas": [{"codigo": "A"}]}],
        }
    ]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)
    titulos = _paragrafos(saida)

    indice_resumo = titulos.index("Resumo")
    indice_importacoes = titulos.index("Importações")
    indice_importacao_1 = titulos.index("Importação 1")
    indice_planilha = titulos.index("Planilha SIGEF")
    assert indice_resumo < indice_importacoes < indice_importacao_1 < indice_planilha


def test_planilha_do_relatorio_nao_traz_pt_original_nem_tipo():
    identificacao = Identificacao()
    parcelas = [{"perimetro": Perimetro(vertices=[Vertice(codigo="A", e_long=-48.0, n_lat=-9.0)]), "importacoes": []}]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)
    doc = Document(io.BytesIO(saida))

    # a última tabela do documento é a Planilha desta (única) parcela
    tabela_planilha = doc.tables[-1]
    cabecalhos = [c.text for c in tabela_planilha.rows[0].cells]
    assert "Pt.Original" not in cabecalhos
    assert "Renum." not in cabecalhos
    assert "Tipo" not in cabecalhos
    assert "Vértice" in cabecalhos


def test_parcela_sem_importacao_avisa_em_vez_de_tabela_vazia():
    identificacao = Identificacao()
    parcelas = [{"perimetro": Perimetro(vertices=[]), "importacoes": []}]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)

    assert "Nenhuma importação nesta parcela." in _paragrafos(saida)


def test_tabelas_tem_fonte_8_cor_preta_e_cabecalho_em_negrito():
    identificacao = Identificacao(nome="Fulano de Tal")
    parcelas = [{"perimetro": Perimetro(vertices=[Vertice(codigo="A", e_long=-48.0, n_lat=-9.0)]), "importacoes": []}]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)
    doc = Document(io.BytesIO(saida))

    for tabela in doc.tables:
        for indice_linha, linha in enumerate(tabela.rows):
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    for run in paragrafo.runs:
                        assert run.font.size == Pt(8)
                        assert run.font.color.rgb == RGBColor(0, 0, 0)
                        assert run.font.bold is (indice_linha == 0)


def test_tabelas_tem_bordas_pretas():
    identificacao = Identificacao()
    parcelas = [{"perimetro": Perimetro(vertices=[Vertice(codigo="A", e_long=-48.0, n_lat=-9.0)]), "importacoes": []}]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)
    doc = Document(io.BytesIO(saida))

    for tabela in doc.tables:
        bordas = tabela._tbl.tblPr.find(qn("w:tblBorders"))
        assert bordas is not None
        for borda in bordas:
            assert borda.get(qn("w:color")) == "000000"


def test_planilha_sigef_fica_em_pagina_paisagem_o_resto_fica_retrato():
    identificacao = Identificacao()
    parcelas = [
        {"perimetro": Perimetro(vertices=[Vertice(codigo="A", e_long=-48.0, n_lat=-9.0)]), "importacoes": []},
        {"perimetro": Perimetro(vertices=[Vertice(codigo="B", e_long=-47.0, n_lat=-8.0)]), "importacoes": []},
    ]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)
    doc = Document(io.BytesIO(saida))

    orientacoes = [secao.orientation for secao in doc.sections]
    # retrato, paisagem (parcela 1), retrato, paisagem (parcela 2) — sem
    # seção extra em branco depois da última paisagem
    assert orientacoes == [
        WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE,
    ]
    secao_paisagem = doc.sections[1]
    assert secao_paisagem.page_width > secao_paisagem.page_height


def test_cabecalho_tem_numero_de_pagina_alinhado_a_direita():
    identificacao = Identificacao()
    saida = gerar_relatorio(identificacao, [{"perimetro": Perimetro(vertices=[]), "importacoes": []}], _CASAS)
    doc = Document(io.BytesIO(saida))

    paragrafo = doc.sections[0].header.paragraphs[0]
    assert paragrafo.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert any(run._r.find(qn("w:fldChar")) is not None for run in paragrafo.runs)


def test_rodape_tem_nome_do_programa_e_data_de_emissao():
    identificacao = Identificacao()
    saida = gerar_relatorio(identificacao, [{"perimetro": Perimetro(vertices=[]), "importacoes": []}], _CASAS)
    doc = Document(io.BytesIO(saida))

    texto_rodape = doc.sections[0].footer.paragraphs[0].text
    assert "Planilha SIGEF" in texto_rodape
    assert datetime.now().strftime("%d/%m/%Y") in texto_rodape


def test_margens_esquerda_2cm_demais_1cm_em_todas_as_secoes():
    identificacao = Identificacao()
    parcelas = [
        {"perimetro": Perimetro(vertices=[Vertice(codigo="A", e_long=-48.0, n_lat=-9.0)]), "importacoes": []},
        {"perimetro": Perimetro(vertices=[Vertice(codigo="B", e_long=-47.0, n_lat=-8.0)]), "importacoes": []},
    ]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)
    doc = Document(io.BytesIO(saida))

    # comparação aproximada: conversão cm -> twips (unidade real do XML de
    # margem) -> EMU não bate byte a byte, só na casa do 0,01 mm.
    assert len(doc.sections) > 1
    for secao in doc.sections:
        assert abs(secao.left_margin - Cm(2)) < Cm(0.01)
        assert abs(secao.right_margin - Cm(1)) < Cm(0.01)
        assert abs(secao.top_margin - Cm(1)) < Cm(0.01)
        assert abs(secao.bottom_margin - Cm(1)) < Cm(0.01)


def test_planilha_sigef_tem_colunas_justas_e_descritivo_mais_largo():
    identificacao = Identificacao()
    parcelas = [{"perimetro": Perimetro(vertices=[Vertice(codigo="A", e_long=-48.0, n_lat=-9.0)]), "importacoes": []}]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)
    doc = Document(io.BytesIO(saida))

    tabela_planilha = doc.tables[-1]
    larguras = [celula.width for celula in tabela_planilha.rows[0].cells]
    cabecalhos = [c.text for c in tabela_planilha.rows[0].cells]

    largura_cns = larguras[cabecalhos.index("CNS")]
    largura_descritivo = larguras[cabecalhos.index("Descritivo")]
    assert largura_cns < largura_descritivo
    assert abs(largura_descritivo - Cm(6.0)) < Cm(0.01)


def test_cabecalho_sigma_long_e_sigma_lat_quebra_em_duas_linhas():
    identificacao = Identificacao()
    parcelas = [{"perimetro": Perimetro(vertices=[Vertice(codigo="A", e_long=-48.0, n_lat=-9.0)]), "importacoes": []}]

    saida = gerar_relatorio(identificacao, parcelas, _CASAS)
    doc = Document(io.BytesIO(saida))

    tabela_planilha = doc.tables[-1]
    cabecalhos = [c.text for c in tabela_planilha.rows[0].cells]
    # python-docx devolve o texto da célula com "\n" onde há quebra de linha
    assert "Sigma\nLong" in cabecalhos
    assert "Sigma\nLat" in cabecalhos
