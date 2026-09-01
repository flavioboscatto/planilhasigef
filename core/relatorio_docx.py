"""Relatório em DOCX: Identificação uma vez no topo e, por parcela (nesta
ordem, pedido explícito do usuário) — resumo do cabeçalho, Importações
(cada sub-aba "Importação N" tal como veio, sem formatação SIGEF) e por
fim a Planilha SIGEF final (só colunas oficiais, mesma formatação de
core/formatacao_vertices.py — sem "Pt.Original" nem "Renum.", que são só
ajuda de tela).

Formatação pedida pelo usuário: fonte 8 e cor preta em toda tabela,
cabeçalho de tabela em negrito, bordas pretas, numeração de página no
topo direito, rodapé com o nome do programa e a data de emissão, e a
tabela da Planilha SIGEF (só ela) em página paisagem — o resto do
relatório fica retrato.

Sem streamlit, sem PySide6 (CLAUDE.md, regra de fronteira de core/): quem
lê os widgets ao vivo pra montar `parcelas` é ui_qt/relatorio_qt.py.
"""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table

from core.formatacao_vertices import COLUNAS_OFICIAIS, ROTULOS_OFICIAIS, linhas_formatadas
from core.models import Identificacao, Perimetro

_PRETO = RGBColor(0, 0, 0)
_TAMANHO_TABELA = Pt(8)

_ROTULOS_IDENTIFICACAO = {
    "natureza_servico": "Natureza do serviço",
    "tipo_pessoa": "Tipo de pessoa",
    "nome": "Nome",
    "cpf_cnpj": "CPF/CNPJ",
    "denominacao_area": "Denominação da área",
    "situacao": "Situação",
    "natureza_area": "Natureza da área",
    "codigo_imovel_sncr": "Código do imóvel (SNCR/INCRA)",
    "codigo_cartorio_cns": "Código do cartório (CNS)",
    "matricula": "Matrícula",
    "municipio_1": "Município",
}


def _pintar_de_preto(run) -> None:
    run.font.color.rgb = _PRETO


def _formatar_celula(celula, tamanho: Pt, negrito: bool = False) -> None:
    for paragrafo in celula.paragraphs:
        for run in paragrafo.runs:
            run.font.size = tamanho
            run.font.bold = negrito
            _pintar_de_preto(run)


def _bordas_pretas(tabela: Table) -> None:
    """python-docx não força cor de borda por estilo de tabela (algumas,
    como "Table Grid", herdam de tema) — desenha a borda preta direto no
    XML, pra nunca depender de qual estilo o Word tem instalado."""
    tbl_pr = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for nome in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borda = OxmlElement(f"w:{nome}")
        borda.set(qn("w:val"), "single")
        borda.set(qn("w:sz"), "4")
        borda.set(qn("w:space"), "0")
        borda.set(qn("w:color"), "000000")
        bordas.append(borda)
    tbl_pr.append(bordas)


def _tabela(doc: Document, cabecalhos: list[str], linhas: list[list[str]]) -> Table | None:
    if not linhas:
        paragrafo = doc.add_paragraph("(nenhum dado)")
        _pintar_de_preto(paragrafo.add_run(""))
        return None
    tabela = doc.add_table(rows=1, cols=len(cabecalhos))
    tabela.style = "Table Grid"
    for celula, texto in zip(tabela.rows[0].cells, cabecalhos):
        celula.text = str(texto)
        _formatar_celula(celula, _TAMANHO_TABELA, negrito=True)
    for linha in linhas:
        celulas = tabela.add_row().cells
        for celula, texto in zip(celulas, linha):
            celula.text = "" if texto is None else str(texto)
            _formatar_celula(celula, _TAMANHO_TABELA)
    _bordas_pretas(tabela)
    return tabela


def _titulo_preto(doc: Document, texto: str, nivel: int):
    cabecalho = doc.add_heading(texto, level=nivel)
    for run in cabecalho.runs:
        _pintar_de_preto(run)
    return cabecalho


def _adicionar_campo_pagina(paragrafo) -> None:
    run = paragrafo.add_run()
    _pintar_de_preto(run)
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instrucao = OxmlElement("w:instrText")
    instrucao.set(qn("xml:space"), "preserve")
    instrucao.text = "PAGE"
    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")
    run._r.append(inicio)
    run._r.append(instrucao)
    run._r.append(fim)


def _configurar_cabecalho_e_rodape(doc: Document) -> None:
    """Só na primeira seção — as seguintes (inclusive as em paisagem da
    Planilha SIGEF) ficam "vinculadas à anterior" por padrão do
    python-docx/Word, então herdam isto automaticamente."""
    secao = doc.sections[0]

    paragrafo_cabecalho = secao.header.paragraphs[0]
    paragrafo_cabecalho.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _adicionar_campo_pagina(paragrafo_cabecalho)

    paragrafo_rodape = secao.footer.paragraphs[0]
    data_emissao = datetime.now().strftime("%d/%m/%Y")
    run_rodape = paragrafo_rodape.add_run(
        f"Relatório emitido pelo programa Planilha SIGEF em {data_emissao}"
    )
    run_rodape.font.size = _TAMANHO_TABELA
    _pintar_de_preto(run_rodape)


def _iniciar_secao_paisagem(doc: Document):
    secao = doc.add_section(WD_SECTION.NEW_PAGE)
    largura, altura = secao.page_height, secao.page_width
    secao.orientation = WD_ORIENT.LANDSCAPE
    secao.page_width = largura
    secao.page_height = altura
    return secao


def _iniciar_secao_retrato(doc: Document):
    secao = doc.add_section(WD_SECTION.NEW_PAGE)
    largura, altura = secao.page_height, secao.page_width
    secao.orientation = WD_ORIENT.PORTRAIT
    secao.page_width = largura
    secao.page_height = altura
    return secao


def gerar_relatorio(identificacao: Identificacao, parcelas: list[dict], casas: dict) -> bytes:
    """`parcelas`: uma entrada por parcela, na ordem das abas — cada uma
    {"perimetro": Perimetro, "importacoes": [{"titulo": str, "colunas":
    list[str], "linhas": list[dict]}, ...]} (ver ui_qt/relatorio_qt.py)."""
    doc = Document()
    _configurar_cabecalho_e_rodape(doc)

    _titulo_preto(doc, "Relatório SIGEF", 1)

    _titulo_preto(doc, "Identificação", 2)
    _tabela(
        doc,
        ["Campo", "Valor"],
        [[rotulo, getattr(identificacao, campo)] for campo, rotulo in _ROTULOS_IDENTIFICACAO.items()],
    )

    for indice, parcela in enumerate(parcelas, start=1):
        perimetro: Perimetro = parcela["perimetro"]
        _titulo_preto(doc, f"Parcela {indice}", 2)

        _titulo_preto(doc, "Resumo", 3)
        _tabela(
            doc,
            ["Campo", "Valor"],
            [
                ["Denominação", perimetro.denominacao],
                ["Lado", perimetro.lado],
                [
                    "Meridiano Central",
                    "" if perimetro.meridiano_central is None else str(perimetro.meridiano_central),
                ],
                ["Hemisfério", perimetro.hemisferio],
            ],
        )

        _titulo_preto(doc, "Importações", 3)
        importacoes = parcela.get("importacoes", [])
        if not importacoes:
            paragrafo = doc.add_paragraph("Nenhuma importação nesta parcela.")
            _pintar_de_preto(paragrafo.runs[0])
        for importacao in importacoes:
            _titulo_preto(doc, importacao["titulo"], 4)
            colunas = importacao["colunas"]
            linhas = [[importacao_linha.get(c, "") for c in colunas] for importacao_linha in importacao["linhas"]]
            _tabela(doc, colunas, linhas)

        # só a Planilha SIGEF vai pra paisagem (fica apertada em retrato) —
        # volta pra retrato logo depois, seja pra próxima parcela ou pro
        # fim do documento.
        _iniciar_secao_paisagem(doc)
        _titulo_preto(doc, "Planilha SIGEF", 3)
        cabecalhos = [ROTULOS_OFICIAIS[c] for c in COLUNAS_OFICIAIS]
        linhas = [[linha[c] for c in COLUNAS_OFICIAIS] for linha in linhas_formatadas(perimetro, casas)]
        _tabela(doc, cabecalhos, linhas)
        # só volta pra retrato se ainda vier outra parcela depois — senão
        # sobra uma página em branco no fim do documento (quebra de seção
        # sempre começa página nova, mesmo sem conteúdo nenhum atrás dela).
        if indice < len(parcelas):
            _iniciar_secao_retrato(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
